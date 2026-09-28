"""Advisory board report (Markdown + DOCX), reusing council.report helpers and
the house style. The deliverable reads like a board memo, not a study."""

from __future__ import annotations

from typing import Any

from council.report import (
    NAVY,
    RED,
    _add_markdown,
    _docx_heading,
    _inline_runs,
    _rated_ideas,
    composites_docx,
    composites_markdown,
)


def to_markdown(result: dict[str, Any]) -> str:
    out: list[str] = []
    out.append("# Advisory board report\n")
    out.append(f"**Question:** {result.get('question', '')}\n")
    out.append(f"**Briefing document used:** {'yes' if result.get('brief_used') else 'no'}\n")

    for w in result.get("warnings") or []:
        out.append(f"> ⚠ {w}\n")

    out.append("## Expert panel\n")
    for e in result.get("experts", []):
        st = e.get("status", "ok")
        detail = f"FAILED — {e.get('error')}" if st == "failed" else st
        out.append(f"- **{e.get('label')}** ({e.get('model')}) — {detail}")
    out.append("")

    out.append("## Recommendation\n")
    out.append(result.get("recommendation", "").strip() or "_(empty)_")
    unsup = result.get("unsupported_claims") or []
    if unsup:
        out.append("\n**⚠ Claims not traceable to any expert:** " + "; ".join(unsup))
    verification = result.get("verification") or []
    if verification:
        out.append("\n**Fact-check:**\n")
        out.append("| Claim | Verdict | Note |")
        out.append("|---|---|---|")
        for v in verification:
            out.append(f"| {v.get('claim','')} | {v.get('verdict','')} | {v.get('note','')} |")
    out.append("")

    out.extend(composites_markdown(result.get("composites") or []))

    evidence = result.get("evidence") or []
    if evidence:
        out.append(f"## Evidence / references\n")
        out.append(f"_Literature grounding — search: {result.get('evidence_query','')}_\n")
        for i, e in enumerate(evidence, 1):
            cite = f"{i}. {e.get('title','')} — {e.get('authors','')}. {e.get('journal','')} {e.get('year','')}."
            if e.get("id"):
                cite += f" PMID:{e['id']}."
            if e.get("url"):
                cite += f" {e['url']}"
            out.append(cite)
        out.append("")

    rated = _rated_ideas(result)
    review = result.get("review") or {}
    if review.get("overall") or rated:
        out.append("## Reviewer assessment\n")
        if review.get("overall"):
            out.append(review["overall"].strip() + "\n")
        if rated:
            out.append("| Point | Verdict | Comment |")
            out.append("|---|---|---|")
            for summary, verdict, a in rated:
                out.append(f"| {summary} | {verdict} | {a.get('comment','')} |")
        out.append("")

    return "\n".join(out)


def to_docx(result: dict[str, Any], path: str) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Garamond"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    r = title.add_run("Advisory Board Report")
    r.bold = True
    r.font.name = "Garamond"
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    sub = doc.add_paragraph()
    sr = sub.add_run(result.get("question", ""))
    sr.italic = True
    sr.font.name = "Garamond"
    doc.add_paragraph(f"Briefing document used: {'yes' if result.get('brief_used') else 'no'}")
    for w in result.get("warnings") or []:
        p = doc.add_paragraph()
        wr = p.add_run("⚠ " + w)
        wr.font.name = "Garamond"
        wr.font.color.rgb = RGBColor.from_string(RED)

    _docx_heading(doc, "Expert panel")
    for e in result.get("experts", []):
        st = e.get("status", "ok")
        detail = f"FAILED — {e.get('error')}" if st == "failed" else st
        p = doc.add_paragraph(style="List Bullet")
        _inline_runs(p, f"**{e.get('label')}** ({e.get('model')}) — {detail}")

    _docx_heading(doc, "Recommendation")
    _add_markdown(doc, result.get("recommendation", "").strip() or "_(empty)_")
    unsup = result.get("unsupported_claims") or []
    if unsup:
        p = doc.add_paragraph()
        rr = p.add_run("⚠ Claims not traceable to any expert: " + "; ".join(unsup))
        rr.font.name = "Garamond"
        rr.font.color.rgb = RGBColor.from_string(RED)
    verification = result.get("verification") or []
    if verification:
        _docx_heading(doc, "Fact-check")
        vt = doc.add_table(rows=1, cols=3)
        vt.style = "Light Grid Accent 1"
        for j, head in enumerate(["Claim", "Verdict", "Note"]):
            vt.rows[0].cells[j].paragraphs[0].add_run(head).bold = True
        for v in verification:
            cells = vt.add_row().cells
            cells[0].text = v.get("claim", "")
            cells[1].text = v.get("verdict", "")
            cells[2].text = v.get("note", "")

    composites_docx(doc, result.get("composites") or [])

    evidence = result.get("evidence") or []
    if evidence:
        _docx_heading(doc, "Evidence / references")
        sub = doc.add_paragraph()
        sr = sub.add_run(f"Literature grounding — search: {result.get('evidence_query','')}")
        sr.italic = True
        sr.font.name = "Garamond"
        for i, e in enumerate(evidence, 1):
            p = doc.add_paragraph(style="List Number")
            cite = f"{e.get('title','')} — {e.get('authors','')}. {e.get('journal','')} {e.get('year','')}."
            if e.get("id"):
                cite += f" PMID:{e['id']}."
            if e.get("url"):
                cite += f" {e['url']}"
            _inline_runs(p, cite)

    rated = _rated_ideas(result)
    review = result.get("review") or {}
    if review.get("overall") or rated:
        _docx_heading(doc, "Reviewer assessment")
        if review.get("overall"):
            _add_markdown(doc, review["overall"].strip())
        if rated:
            rt = doc.add_table(rows=1, cols=3)
            rt.style = "Light Grid Accent 1"
            for j, head in enumerate(["Point", "Verdict", "Comment"]):
                rt.rows[0].cells[j].paragraphs[0].add_run(head).bold = True
            for summary, verdict, a in rated:
                cells = rt.add_row().cells
                cells[0].text = summary
                cells[1].text = verdict
                cells[2].text = a.get("comment", "")

    doc.save(path)
